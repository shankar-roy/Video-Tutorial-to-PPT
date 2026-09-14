#!/usr/bin/env python3
"""Render a design.json spec into a .pptx, animations included.

The spec is authored by reading tutorial frames; this script only turns it into
OOXML. See design-reference.md for the schema.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("Install python-pptx: python3 -m pip install python-pptx")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_anim import add_transition, attach_timing  # noqa: E402

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "round1Rect": MSO_SHAPE.ROUND_1_RECTANGLE,
    "round2SameRect": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "pie": MSO_SHAPE.PIE,
    "chord": MSO_SHAPE.CHORD,
    "arc": MSO_SHAPE.ARC,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "rightTriangle": MSO_SHAPE.RIGHT_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "star5": MSO_SHAPE.STAR_5_POINT,
    "star4": MSO_SHAPE.STAR_4_POINT,
    "arrowRight": MSO_SHAPE.RIGHT_ARROW,
    "chevron": MSO_SHAPE.CHEVRON,
    "teardrop": MSO_SHAPE.TEAR,
    "donut": MSO_SHAPE.DONUT,
    "blockArc": MSO_SHAPE.BLOCK_ARC,
    "plaque": MSO_SHAPE.PLAQUE,
    "line": MSO_SHAPE.RECTANGLE,  # thin rect: crisper than a connector
}

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(str(value).lstrip("#").upper())


def sub(parent, tag: str):
    from lxml import etree

    return etree.SubElement(parent, qn(tag))


def add_shadow(effects, spec) -> None:
    inner = str(spec.get("type", "outer")).lower() == "inner"
    shadow = sub(effects, "a:innerShdw" if inner else "a:outerShdw")
    shadow.set("blurRad", str(Pt(spec.get("blur", 12)).emu))
    shadow.set("dist", str(Pt(spec.get("dist", 4)).emu))
    shadow.set("dir", str(spec.get("dir", 5400000)))
    if not inner:
        shadow.set("rotWithShape", "0")
    color = sub(shadow, "a:srgbClr")
    color.set("val", str(spec.get("color", "000000")).lstrip("#").upper())
    alpha = sub(color, "a:alpha")
    alpha.set("val", str(int(float(spec.get("alpha", 0.25)) * 100000)))


def set_shadow(shape, spec) -> None:
    """False/None clears shadows; a dict or a list of dicts sets them.

    a:effectLst holds at most one innerShdw and one outerShdw, in that order, so
    a neumorphic pair of outer shadows has to come from two stacked shapes.
    """
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    effects = sub(spPr, "a:effectLst")
    if not spec:
        return
    specs = spec if isinstance(spec, list) else [spec]
    for one in sorted(specs, key=lambda s: str(s.get("type", "outer")).lower() != "inner"):
        add_shadow(effects, one)


def set_fill(shape, spec) -> None:
    if spec is None:
        shape.fill.background()
        return
    if isinstance(spec, str):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(spec)
        return
    if spec.get("type") == "gradient":
        stops = spec.get("stops", [["FFFFFF", 0], ["CCCCCC", 100]])
        shape.fill.gradient()
        shape.fill.gradient_angle = spec.get("angle", 90)
        existing = shape.fill.gradient_stops
        for i, (color, pos) in enumerate(stops[: len(existing)]):
            existing[i].color.rgb = rgb(color)
            existing[i].position = float(pos) / 100.0
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(spec.get("color", "FFFFFF"))


def set_line(shape, spec) -> None:
    if not spec:
        shape.line.fill.background()
        return
    shape.line.color.rgb = rgb(spec.get("color", "000000"))
    shape.line.width = Pt(spec.get("width", 1))
    cap = spec.get("cap")
    if cap:
        shape.line._get_or_add_ln().set("cap", cap)
    if spec.get("dash"):
        ln = shape.line._get_or_add_ln()
        for old in ln.findall(qn("a:prstDash")):
            ln.remove(old)
        sub(ln, "a:prstDash").set("val", spec["dash"])
    if spec.get("round_ends"):
        ln = shape.line._get_or_add_ln()
        ln.set("cap", "rnd")
        sub(ln, "a:round")


def set_text(shape, spec: dict, defaults: dict) -> None:
    text = spec.get("text")
    runs = spec.get("runs")
    if text is None and not runs:
        return

    frame = shape.text_frame
    frame.word_wrap = spec.get("wrap", True)
    for side in ("left", "right", "top", "bottom"):
        frame.__setattr__(f"margin_{side}", Inches(spec.get(f"margin_{side}", 0.02)))
    font_spec = {**defaults, **(spec.get("font") or {})}
    frame.vertical_anchor = ANCHOR.get(font_spec.get("anchor", "middle"), MSO_ANCHOR.MIDDLE)

    blocks = runs if runs else [{"text": text, **font_spec}]
    lines = []
    for block in blocks:
        for piece in str(block.get("text", "")).split("\n"):
            lines.append({**font_spec, **block, "text": piece})

    frame.clear()
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = ALIGN.get(line.get("align", "center"), PP_ALIGN.CENTER)
        if line.get("line_spacing"):
            para.line_spacing = float(line["line_spacing"])
        if line.get("space_after") is not None:
            para.space_after = Pt(line["space_after"])
        run = para.add_run()
        run.text = line["text"]
        font = run.font
        font.name = line.get("name", "Calibri")
        font.size = Pt(line.get("size", 18))
        font.bold = bool(line.get("bold", False))
        font.italic = bool(line.get("italic", False))
        font.color.rgb = rgb(line.get("color", "000000"))
        if line.get("spacing"):
            # character spacing, in 1/100 pt
            run.font._rPr.set("spc", str(int(float(line["spacing"]) * 100)))
        if line.get("caps"):
            run.font._rPr.set("cap", "all")
        latin = run.font._rPr.find(qn("a:latin"))
        if latin is None:
            latin = sub(run.font._rPr, "a:latin")
        latin.set("typeface", line.get("name", "Calibri"))


def add_freeform(slide, spec: dict):
    points = spec.get("points") or []
    if len(points) < 2:
        raise ValueError(f"freeform '{spec.get('id')}' needs at least 2 points")
    builder = slide.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]))
    builder.add_line_segments(
        [(Inches(x), Inches(y)) for x, y in points[1:]], close=spec.get("close", True)
    )
    return builder.convert_to_shape()


def _set_shape_name(shape, spec: dict) -> None:
    """Morph pairs objects across slides by this name. Prefer an explicit
    `name`; fall back to `id` so repeated year tabs stay matched."""
    name = spec.get("name") or spec.get("id")
    if name:
        shape.name = str(name)


def make_shape(slide, spec: dict, defaults: dict):
    kind = spec.get("type", "rect")

    if kind == "freeform":
        shape = add_freeform(slide, spec)
    elif kind == "text":
        shape = slide.shapes.add_textbox(
            Inches(spec["x"]), Inches(spec["y"]), Inches(spec["w"]), Inches(spec["h"])
        )
    elif kind == "picture":
        shape = slide.shapes.add_picture(
            spec["path"], Inches(spec["x"]), Inches(spec["y"]),
            Inches(spec["w"]), Inches(spec["h"]),
        )
        _set_shape_name(shape, spec)
        return shape
    else:
        if kind not in SHAPES:
            raise ValueError(f"Unknown shape type: {kind}")
        shape = slide.shapes.add_shape(
            SHAPES[kind], Inches(spec["x"]), Inches(spec["y"]),
            Inches(spec["w"]), Inches(spec["h"]),
        )

    _set_shape_name(shape, spec)

    if kind != "text":
        set_fill(shape, spec.get("fill", "FFFFFF"))
        set_line(shape, spec.get("line"))
        set_shadow(shape, spec.get("shadow", False))
        adj = spec.get("adj")
        if adj is not None:
            values = adj if isinstance(adj, (list, tuple)) else [adj]
            for i, value in enumerate(values):
                if i < len(shape.adjustments):
                    shape.adjustments[i] = float(value)
    else:
        shape.fill.background()
        shape.line.fill.background()

    if spec.get("rotation"):
        shape.rotation = float(spec["rotation"])
    if spec.get("flip_h"):
        shape._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
    if spec.get("flip_v"):
        shape._element.spPr.find(qn("a:xfrm")).set("flipV", "1")

    set_text(shape, spec, defaults)
    return shape


def set_background(slide, prs, spec) -> None:
    """Write a real <p:bg>: solid, linear gradient, or radial gradient."""
    if not spec:
        return
    c_sld = slide._element.cSld
    for old in c_sld.findall(qn("p:bg")):
        c_sld.remove(old)
    bg = c_sld.makeelement(qn("p:bg"), {})
    c_sld.insert(0, bg)
    bg_pr = sub(bg, "p:bgPr")

    if isinstance(spec, str) or spec.get("type") == "solid":
        color = spec if isinstance(spec, str) else spec.get("color", "FFFFFF")
        sub(sub(bg_pr, "a:solidFill"), "a:srgbClr").set(
            "val", str(color).lstrip("#").upper()
        )
        sub(bg_pr, "a:effectLst")
        return

    grad = sub(bg_pr, "a:gradFill")
    grad.set("flip", "none")
    grad.set("rotWithShape", "1")
    gs_lst = sub(grad, "a:gsLst")
    for color, pos in spec.get("stops", [["FFFFFF", 0], ["CCCCCC", 100]]):
        gs = sub(gs_lst, "a:gs")
        gs.set("pos", str(int(float(pos) * 1000)))
        sub(gs, "a:srgbClr").set("val", str(color).lstrip("#").upper())

    if spec.get("type") == "radial":
        path = sub(grad, "a:path")
        path.set("path", "circle")
        center = spec.get("center", [50, 50])
        rect = sub(path, "a:fillToRect")
        rect.set("l", str(int(center[0] * 1000)))
        rect.set("t", str(int(center[1] * 1000)))
        rect.set("r", str(int((100 - center[0]) * 1000)))
        rect.set("b", str(int((100 - center[1]) * 1000)))
    else:
        lin = sub(grad, "a:lin")
        lin.set("ang", str(int(spec.get("angle", 90) * 60000)))
        lin.set("scaled", "0")
    sub(bg_pr, "a:effectLst")


def normalize_groups(animations) -> list:
    """Accept either a flat step list or explicit click groups."""
    if not animations:
        return []
    if isinstance(animations, dict):
        animations = [animations]
    if all(isinstance(a, dict) and "steps" in a for a in animations):
        return [a["steps"] for a in animations]
    return [animations]


def expand_targets(steps: list, ids: dict) -> list:
    """A step may name several targets; they animate together."""
    out = []
    for step in steps:
        targets = step.get("targets") or [step.get("target")]
        for i, name in enumerate(targets):
            if name not in ids:
                raise ValueError(f"animation targets unknown shape id: {name!r}")
            row = {k: v for k, v in step.items() if k not in {"target", "targets"}}
            row["spid"] = ids[name]
            if i:
                row["start"] = "with_previous"
                row["delay"] = 0
            out.append(row)
    return out


def build_slide(prs, slide_spec: dict, defaults: dict) -> int:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, prs, slide_spec.get("bg"))

    ids = {}
    for spec in slide_spec.get("shapes", []):
        shape = make_shape(slide, spec, defaults)
        if spec.get("id"):
            ids[spec["id"]] = shape.shape_id

    groups = [expand_targets(steps, ids) for steps in normalize_groups(slide_spec.get("animations"))]
    count = attach_timing(slide, groups)
    kind, speed = _transition_args(slide_spec)
    if kind:
        add_transition(slide, kind, speed)
    return count


def _transition_args(slide_spec: dict) -> tuple:
    transition = slide_spec.get("transition")
    speed = slide_spec.get("transition_speed")
    if isinstance(transition, dict):
        kind = transition.get("kind") or transition.get("type")
        speed = transition.get("speed", speed)
    else:
        kind = transition
    if not kind:
        return None, None
    if not speed:
        speed = "slow" if str(kind).lower() == "morph" else "med"
    return kind, speed


def build(spec: dict, dest: Path) -> dict:
    prs = Presentation()
    canvas = spec.get("canvas", {})
    prs.slide_width = Inches(canvas.get("w_in", 13.333))
    prs.slide_height = Inches(canvas.get("h_in", 7.5))

    defaults = {"name": "Calibri", "size": 18, "color": "000000", **(spec.get("defaults") or {})}
    slides = spec.get("slides")
    if slides is None:
        slides = [{
            "shapes": spec.get("shapes", []),
            "animations": spec.get("animations"),
            "bg": canvas.get("bg"),
            "transition": spec.get("transition"),
            "transition_speed": spec.get("transition_speed"),
        }]

    total = 0
    for slide_spec in slides:
        slide_spec.setdefault("bg", canvas.get("bg"))
        total += build_slide(prs, slide_spec, defaults)

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    return {"path": str(dest), "slides": len(slides), "animation_steps": total}


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a .pptx from a design.json spec.")
    parser.add_argument("spec", help="Path to design.json")
    parser.add_argument("out", help="Destination .pptx")
    args = parser.parse_args()

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    result = build(spec, Path(args.out))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
