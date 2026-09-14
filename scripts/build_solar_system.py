#!/usr/bin/env python3
"""Infographic solar-system slide matching the PowerPoint School target layout."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt, Emu
except ImportError:
    sys.exit("Install python-pptx: python3 -m pip install python-pptx")

SLIDE_W = 13.333
SLIDE_H = 7.5
CX = SLIDE_W / 2
CY = SLIDE_H / 2

# Positions are inches on a 13.333" × 7.5" slide, traced from the target infographic.
BODIES = [
    {
        "name": "NEPTUNE",
        "blurb": "THE ENIGMATIC BLUE\nPLANET WITH VERY COLD\nTEMPERATURE",
        "fill": "4A90D9",
        "shade": "2E5F9E",
        "title": "4AA4E8",
        "size": 0.82,
        "x": 0.22,
        "y": 0.18,
        "label_x": 1.18,
        "label_y": 0.18,
        "label_w": 3.35,
        "align": "left",
        "shade_rot": 200,
    },
    {
        "name": "SATURN",
        "blurb": "GAS GIANT & SECOND LARGEST\nPLANET AFTER JUPITER IN OUR\nSOLAR SYSTEM",
        "fill": "C45BB0",
        "shade": "8A3A88",
        "title": "C85BB8",
        "size": 0.72,
        "x": 12.28,
        "y": 0.28,
        "label_x": 8.55,
        "label_y": 0.18,
        "label_w": 3.55,
        "align": "right",
        "shade_rot": 40,
    },
    {
        "name": "MERCURY",
        "blurb": "CLOSEST PLANET TO THE\nSUN WITH HIGH TEMPERATURE\n& A VERY THIN ATMOSPHERE",
        "fill": "E0B070",
        "shade": "C48A48",
        "title": "E0A85C",
        "size": 0.58,
        "x": 6.22,
        "y": 1.72,
        "label_x": 2.85,
        "label_y": 1.35,
        "label_w": 3.25,
        "align": "right",
        "shade_rot": 30,
    },
    {
        "name": "MOON",
        "blurb": "RELATIVELY SMALL\n& ONE FOURTH THE\nSIZE OF EARTH",
        "fill": "D8DCE8",
        "shade": "9AA3B8",
        "title": "E8EEF8",
        "size": 0.38,
        "x": 9.72,
        "y": 2.22,
        "label_x": 6.95,
        "label_y": 1.85,
        "label_w": 2.6,
        "align": "right",
        "shade_rot": 200,
    },
    {
        "name": "EARTH",
        "blurb": "ONLY KNOWN PLANET TO\nSUPPORT LIFE & ITS DIVERSE\nECOSYSTEMS",
        "fill": "4DB8B0",
        "shade": "2E8A86",
        "title": "4DC8C0",
        "size": 0.68,
        "x": 10.05,
        "y": 3.42,
        "label_x": 10.88,
        "label_y": 3.35,
        "label_w": 2.35,
        "align": "left",
        "shade_rot": 200,
    },
    {
        "name": "MARS",
        "blurb": "EVIDENCE SUGGESTS\nTHAT LIQUID WATER MAY\nEXIST THERE",
        "fill": "E07030",
        "shade": "B84818",
        "title": "E06828",
        "size": 0.58,
        "x": 2.18,
        "y": 3.28,
        "label_x": 0.15,
        "label_y": 3.05,
        "label_w": 1.95,
        "align": "right",
        "shade_rot": 200,
    },
    {
        "name": "JUPITER",
        "blurb": "GAS GIANT CONSISTS\nOF HYDROGEN & HELIUM\nWITH NO SOLID SURFACE",
        "fill": "E070B0",
        "shade": "B84888",
        "title": "E070B8",
        "size": 0.78,
        "x": 2.28,
        "y": 5.28,
        "label_x": 3.18,
        "label_y": 5.22,
        "label_w": 2.9,
        "align": "left",
        "shade_rot": 200,
    },
    {
        "name": "VENUS",
        "blurb": "THICK ATMOSPHERE\nMOSTLY COMPOSED OF\nCARBON DIOXIDE",
        "fill": "D08040",
        "shade": "A85828",
        "title": "E08040",
        "size": 0.62,
        "x": 5.18,
        "y": 6.28,
        "label_x": 5.95,
        "label_y": 6.12,
        "label_w": 2.55,
        "align": "left",
        "shade_rot": 200,
    },
    {
        "name": "URANUS",
        "blurb": "ICE GIANT ROTATES ON\nITS SIDE WITH DISTINCT\nAPPEARANCE",
        "fill": "5EC8E0",
        "shade": "2E98B0",
        "title": "4EC0E0",
        "size": 0.78,
        "x": 12.22,
        "y": 5.88,
        "label_x": 8.85,
        "label_y": 5.85,
        "label_w": 3.2,
        "align": "right",
        "shade_rot": 200,
    },
]

STARS = [
    (1.55, 1.85, 0.10),
    (2.85, 0.55, 0.08),
    (3.55, 2.55, 0.07),
    (4.35, 0.95, 0.09),
    (4.85, 4.55, 0.08),
    (5.55, 2.15, 0.07),
    (7.85, 0.65, 0.09),
    (8.15, 5.45, 0.08),
    (8.95, 4.15, 0.07),
    (9.55, 6.55, 0.08),
    (11.15, 2.05, 0.09),
    (11.55, 4.85, 0.07),
    (12.65, 1.65, 0.08),
    (0.95, 5.15, 0.07),
    (1.85, 6.65, 0.09),
    (6.85, 6.85, 0.07),
    (10.35, 0.45, 0.08),
    (3.15, 6.15, 0.07),
]


def rgb(hex6: str) -> RGBColor:
    hex6 = hex6.lstrip("#")
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


def no_shadow(shape) -> None:
    shape.shadow.inherit = False


def set_italic_run(paragraph, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    paragraph.clear() if hasattr(paragraph, "clear") else None
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.italic = True
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.font.name = "Arial Narrow"
    rpr = paragraph.runs[0]._r.get_or_add_rPr()
    spc = rpr.find(qn("a:spc"))
    if spc is None:
        spc = etree.SubElement(rpr, qn("a:spc"))
    spc.set("val", "80")


def set_radial_bg(slide) -> None:
    c_sld = slide._element.cSld
    bg = c_sld.find(qn("p:bg"))
    if bg is not None:
        c_sld.remove(bg)
    bg = etree.SubElement(c_sld, qn("p:bg"))
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    grad = etree.SubElement(bg_pr, qn("a:gradFill"))
    grad.set("flip", "none")
    grad.set("rotWithShape", "1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, color in (("0", "2A3F8A"), ("40000", "1A2268"), ("100000", "0A071C")):
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", pos)
        c = etree.SubElement(gs, qn("a:srgbClr"))
        c.set("val", color)
    path = etree.SubElement(grad, qn("a:path"))
    path.set("path", "circle")
    fill_to = etree.SubElement(path, qn("a:fillToRect"))
    fill_to.set("l", "50000")
    fill_to.set("t", "50000")
    fill_to.set("r", "50000")
    fill_to.set("b", "50000")
    etree.SubElement(bg_pr, qn("a:effectLst"))


PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def add_orbits(slide):
    shapes = []
    for i, size in enumerate((2.55, 3.55, 4.70, 5.90, 7.15, 8.45, 9.80, 11.20, 12.70, 14.20), 1):
        sh = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(CX - size / 2),
            Inches(CY - size / 2),
            Inches(size),
            Inches(size),
        )
        no_shadow(sh)
        sh.fill.background()
        sh.line.color.rgb = rgb("7B8AB8")
        sh.line.width = Pt(1.25)
        sh.line.dash_style = MSO_LINE.ROUND_DOT
        sh.name = f"Orbit {i}"
        shapes.append(sh)
    return shapes


def add_sun(slide):
    shapes = []
    rings = ((2.46, "3D2810"), (2.28, "5A3A18"), (1.98, "E09020"), (1.62, "FFC400"))
    for i, (size, color) in enumerate(rings, 1):
        sh = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(CX - size / 2),
            Inches(CY - size / 2),
            Inches(size),
            Inches(size),
        )
        no_shadow(sh)
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()
        sh.name = f"Sun {i}"
        shapes.append(sh)
    return shapes


def set_pie_180(shape) -> None:
    prst = shape._element.spPr.find(qn("a:prstGeom"))
    av_lst = prst.find(qn("a:avLst"))
    if av_lst is None:
        av_lst = etree.SubElement(prst, qn("a:avLst"))
    else:
        for child in list(av_lst):
            av_lst.remove(child)
    for name, val in (("adj1", "0"), ("adj2", "10800000")):
        gd = etree.SubElement(av_lst, qn("a:gd"))
        gd.set("name", name)
        gd.set("fmla", f"val {val}")


def add_planet(slide, body: dict) -> dict:
    x, y, size = body["x"], body["y"], body["size"]
    base = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    no_shadow(base)
    base.fill.solid()
    base.fill.fore_color.rgb = rgb(body["fill"])
    base.line.fill.background()
    base.name = f"{body['name']} body"

    shade = slide.shapes.add_shape(MSO_SHAPE.PIE, Inches(x), Inches(y), Inches(size), Inches(size))
    no_shadow(shade)
    shade.fill.solid()
    shade.fill.fore_color.rgb = rgb(body["shade"])
    shade.line.fill.background()
    set_pie_180(shade)
    shade.rotation = body["shade_rot"]
    shade.name = f"{body['name']} shade"

    box = slide.shapes.add_textbox(
        Inches(body["label_x"]),
        Inches(body["label_y"]),
        Inches(body["label_w"]),
        Inches(1.05),
    )
    box.name = f"{body['name']} label"
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    align = PP_ALIGN.LEFT if body["align"] == "left" else PP_ALIGN.RIGHT
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = body["name"]
    p.font.size = Pt(18)
    p.font.italic = True
    p.font.bold = True
    p.font.color.rgb = rgb(body["title"])
    p.font.name = "Arial Narrow"

    for line in body["blurb"].split("\n"):
        para = tf.add_paragraph()
        para.alignment = align
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        para.text = line
        para.font.size = Pt(10)
        para.font.italic = True
        para.font.bold = True
        para.font.color.rgb = RGBColor(0xF4, 0xF6, 0xFB)
        para.font.name = "Arial Narrow"
    return {"name": body["name"], "align": body["align"], "body": base, "shade": shade, "label": box}


def add_stars(slide):
    shapes = []
    for i, (x, y, s) in enumerate(STARS, 1):
        star = slide.shapes.add_shape(
            MSO_SHAPE.STAR_4_POINT, Inches(x), Inches(y), Inches(s), Inches(s)
        )
        no_shadow(star)
        star.fill.solid()
        star.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFF)
        star.line.fill.background()
        star.name = f"Star {i}"
        shapes.append(star)
    return shapes


class Ids:
    def __init__(self) -> None:
        self.n = 2

    def next(self) -> str:
        self.n += 1
        return str(self.n)


def _e(tag, attrib=None, children=None, text=None):
    ns = PML if not tag.startswith("{") else ""
    el = etree.Element((ns + tag) if ns else tag, attrib or {})
    if text is not None:
        el.text = text
    for child in children or []:
        if child is not None:
            el.append(child)
    return el


def tgt(spid) -> etree._Element:
    return _e("tgtEl", children=[_e("spTgt", {"spid": str(spid)})])


def cond_delay(delay) -> etree._Element:
    return _e("stCondLst", children=[_e("cond", {"delay": str(delay)})])


def set_visible(ids: Ids, spid) -> etree._Element:
    return _e(
        "set",
        children=[
            _e(
                "cBhvr",
                children=[
                    _e("cTn", {"id": ids.next(), "dur": "1", "fill": "hold"}, [cond_delay(0)]),
                    tgt(spid),
                    _e("attrNameLst", children=[_e("attrName", text="style.visibility")]),
                ],
            ),
            _e("to", children=[_e("strVal", {"val": "visible"})]),
        ],
    )


def zoom_behaviors(ids: Ids, spid, dur: int) -> list:
    def size_anim(attr: str) -> etree._Element:
        return _e(
            "anim",
            {"calcmode": "lin", "valueType": "num"},
            [
                _e(
                    "cBhvr",
                    children=[
                        _e("cTn", {"id": ids.next(), "dur": str(dur), "fill": "hold"}),
                        tgt(spid),
                        _e("attrNameLst", children=[_e("attrName", text=attr)]),
                    ],
                ),
                _e(
                    "tavLst",
                    children=[
                        _e("tav", {"tm": "0"}, [_e("val", children=[_e("fltVal", {"val": "0"})])]),
                        _e(
                            "tav",
                            {"tm": "100000"},
                            [_e("val", children=[_e("strVal", {"val": f"#{attr}"})])],
                        ),
                    ],
                ),
            ],
        )

    fade = _e(
        "animEffect",
        {"transition": "in", "filter": "fade"},
        [_e("cBhvr", children=[_e("cTn", {"id": ids.next(), "dur": str(dur)}), tgt(spid)])],
    )
    return [set_visible(ids, spid), size_anim("ppt_w"), size_anim("ppt_h"), fade]


def spinner_behaviors(ids: Ids, spid, dur: int) -> list:
    rot = _e(
        "animRot",
        {"by": "21600000"},
        [
            _e(
                "cBhvr",
                children=[
                    _e("cTn", {"id": ids.next(), "dur": str(dur), "fill": "hold"}),
                    tgt(spid),
                    _e("attrNameLst", children=[_e("attrName", text="r")]),
                ],
            )
        ],
    )
    fade = _e(
        "animEffect",
        {"transition": "in", "filter": "fade"},
        [_e("cBhvr", children=[_e("cTn", {"id": ids.next(), "dur": str(dur)}), tgt(spid)])],
    )
    return [set_visible(ids, spid), rot, fade]


def grow_shrink(ids: Ids, spid, dur: int = 2000) -> etree._Element:
    return _e(
        "animScale",
        children=[
            _e(
                "cBhvr",
                children=[
                    _e(
                        "cTn",
                        {
                            "id": ids.next(),
                            "dur": str(dur),
                            "fill": "hold",
                            "autoRev": "1",
                            "repeatCount": "indefinite",
                            "accel": "50000",
                            "decel": "50000",
                        },
                    ),
                    tgt(spid),
                ],
            ),
            _e("to", children=[_e("ptStr", {"x": "1.08", "y": "1.08"})]),
        ],
    )


def spin_emphasis(ids: Ids, spid, dur: int, clockwise: bool) -> etree._Element:
    by = "21600000" if clockwise else "-21600000"
    return _e(
        "animRot",
        {"by": by},
        [
            _e(
                "cBhvr",
                children=[
                    _e(
                        "cTn",
                        {
                            "id": ids.next(),
                            "dur": str(dur),
                            "fill": "hold",
                            "repeatCount": "indefinite",
                        },
                    ),
                    tgt(spid),
                    _e("attrNameLst", children=[_e("attrName", text="r")]),
                ],
            )
        ],
    )


def fade_in(ids: Ids, spid, dur: int) -> list:
    fade = _e(
        "animEffect",
        {"transition": "in", "filter": "fade"},
        [_e("cBhvr", children=[_e("cTn", {"id": ids.next(), "dur": str(dur)}), tgt(spid)])],
    )
    return [set_visible(ids, spid), fade]


def motion_line(ids: Ids, spid, dur: int, left: bool) -> etree._Element:
    x = "-0.08" if left else "0.08"
    return _e(
        "animMotion",
        {"origin": "layout", "path": f"M 0 0 L {x} 0 E", "pathEditMode": "relative"},
        [
            _e(
                "cBhvr",
                children=[
                    _e(
                        "cTn",
                        {
                            "id": ids.next(),
                            "dur": str(dur),
                            "fill": "hold",
                            "autoRev": "1",
                            "repeatCount": "indefinite",
                        },
                    ),
                    tgt(spid),
                    _e(
                        "attrNameLst",
                        children=[_e("attrName", text="ppt_x"), _e("attrName", text="ppt_y")],
                    ),
                ],
            )
        ],
    )


def color_pulse(ids: Ids, spid, dur: int = 2500) -> etree._Element:
    to = _e("to")
    clr = etree.SubElement(to, f"{DML}srgbClr")
    clr.set("val", "FFE566")
    return _e(
        "animClr",
        {"clrSpc": "rgb"},
        [
            _e(
                "cBhvr",
                children=[
                    _e(
                        "cTn",
                        {
                            "id": ids.next(),
                            "dur": str(dur),
                            "fill": "hold",
                            "autoRev": "1",
                            "repeatCount": "indefinite",
                        },
                    ),
                    tgt(spid),
                    _e("attrNameLst", children=[_e("attrName", text="fillcolor")]),
                ],
            ),
            to,
        ],
    )


def effect_par(ids: Ids, preset_id, preset_class, subtype, node_type, behaviors) -> etree._Element:
    return _e(
        "par",
        children=[
            _e(
                "cTn",
                {
                    "id": ids.next(),
                    "presetID": str(preset_id),
                    "presetClass": preset_class,
                    "presetSubtype": str(subtype),
                    "fill": "hold",
                    "grpId": "0",
                    "nodeType": node_type,
                },
                [cond_delay(0), _e("childTnLst", children=behaviors)],
            )
        ],
    )


def delayed_par(ids: Ids, delay_ms: int, children) -> etree._Element:
    return _e(
        "par",
        children=[
            _e(
                "cTn",
                {"id": ids.next(), "fill": "hold"},
                [cond_delay(delay_ms), _e("childTnLst", children=children)],
            )
        ],
    )


def click_group(ids: Ids, children) -> etree._Element:
    inner = _e(
        "par",
        children=[
            _e(
                "cTn",
                {"id": ids.next(), "fill": "hold"},
                [cond_delay(0), _e("childTnLst", children=children)],
            )
        ],
    )
    return _e(
        "par",
        children=[
            _e(
                "cTn",
                {"id": ids.next(), "fill": "hold"},
                [cond_delay("indefinite"), _e("childTnLst", children=[inner])],
            )
        ],
    )


def attach_timing(slide, orbits, suns, stars, planets) -> None:
    ids = Ids()
    click_pars = []

    first_bits = []
    for i, sh in enumerate(orbits):
        node = "clickEffect" if i == 0 else "withEffect"
        first_bits.append(
            delayed_par(
                ids,
                0 if i == 0 else i * 50,
                [
                    effect_par(
                        ids, 23, "entr", 0, node, zoom_behaviors(ids, sh.shape_id, 750)
                    )
                ],
            )
        )
    for i, sh in enumerate(suns):
        first_bits.append(
            delayed_par(
                ids,
                i * 150,
                [
                    effect_par(
                        ids, 44, "entr", 0, "withEffect", spinner_behaviors(ids, sh.shape_id, 1000)
                    ),
                    effect_par(ids, 6, "emph", 0, "withEffect", [grow_shrink(ids, sh.shape_id, 2000)]),
                ],
            )
        )
    for sh in stars:
        first_bits.append(
            delayed_par(
                ids,
                0,
                [
                    effect_par(ids, 10, "entr", 0, "withEffect", fade_in(ids, sh.shape_id, 250)),
                    effect_par(ids, 9, "emph", 0, "withEffect", [color_pulse(ids, sh.shape_id)]),
                ],
            )
        )
    click_pars.append(click_group(ids, first_bits))

    order = ["MERCURY", "VENUS", "EARTH", "MOON", "MARS", "JUPITER", "SATURN", "URANUS", "NEPTUNE"]
    by_name = {p["name"]: p for p in planets}
    for i, name in enumerate(order):
        p = by_name[name]
        clockwise = i % 2 == 1
        bits = []
        for sh in (p["body"], p["shade"]):
            bits.append(
                effect_par(
                    ids,
                    8,
                    "emph",
                    0,
                    "clickEffect" if sh is p["body"] else "withEffect",
                    [spin_emphasis(ids, sh.shape_id, 15000, clockwise)],
                )
            )
        label_id = p["label"].shape_id
        left = p["align"] == "left"
        bits.append(
            delayed_par(
                ids,
                750,
                [
                    effect_par(ids, 10, "entr", 0, "withEffect", fade_in(ids, label_id, 250)),
                    effect_par(
                        ids,
                        0,
                        "path",
                        0,
                        "withEffect",
                        [motion_line(ids, label_id, 2000, left=left)],
                    ),
                ],
            )
        )
        click_pars.append(click_group(ids, bits))

    seq = _e(
        "seq",
        {"concurrent": "1", "nextAc": "seek"},
        [
            _e(
                "cTn",
                {"id": "2", "dur": "indefinite", "nodeType": "mainSeq"},
                [_e("childTnLst", children=click_pars)],
            ),
            _e(
                "prevCondLst",
                children=[
                    _e("cond", {"evt": "onPrev", "delay": "0"}, [ _e("tgtEl", children=[_e("sldTgt")])])
                ],
            ),
            _e(
                "nextCondLst",
                children=[
                    _e("cond", {"evt": "onNext", "delay": "0"}, [_e("tgtEl", children=[_e("sldTgt")])])
                ],
            ),
        ],
    )
    root = _e(
        "timing",
        children=[
            _e(
                "tnLst",
                children=[
                    _e(
                        "par",
                        children=[
                            _e(
                                "cTn",
                                {
                                    "id": "1",
                                    "dur": "indefinite",
                                    "restart": "never",
                                    "nodeType": "tmRoot",
                                },
                                [_e("childTnLst", children=[seq])],
                            )
                        ],
                    )
                ],
            )
        ],
    )
    seen = []
    for coll in (orbits, suns, stars):
        seen.extend(coll)
    for p in planets:
        seen.extend([p["body"], p["shade"], p["label"]])
    bld = _e("bldLst")
    used = set()
    for sh in seen:
        spid = str(sh.shape_id)
        if spid in used:
            continue
        used.add(spid)
        bld.append(_e("bldP", {"spid": spid, "grpId": "0", "anim": "1"}))
    root.append(bld)
    slide._element.append(root)


def build(dest: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_radial_bg(slide)
    orbits = add_orbits(slide)
    stars = add_stars(slide)
    suns = add_sun(slide)
    planets = [add_planet(slide, body) for body in BODIES]
    attach_timing(slide, orbits, suns, stars, planets)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    print(dest)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("slides/solar-system.pptx")
    build(out)
