#!/usr/bin/env python3
"""Build a 16:9 PowerPoint file from slides.json."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    sys.exit("Install python-pptx: python3 -m pip install python-pptx")

BG = RGBColor(0x12, 0x18, 0x24)
TITLE_COLOR = RGBColor(0xF8, 0xFA, 0xFC)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)
BODY = RGBColor(0xE2, 0xE8, 0xF0)
MUTED = RGBColor(0x94, 0xA3, 0xB8)

VALID_LAYOUTS = {"title", "section", "bullets"}
VALID_INTENTS = {"follow_instructions", "teaching_outline", "mixed"}


def style_paragraph(paragraph, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    paragraph.font.name = "Calibri"


def fill_slide(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def notes_for(slide, text: str | None) -> None:
    if not text:
        return
    slide.notes_slide.notes_text_frame.text = text


def render_title(slide, item: dict) -> None:
    tf = add_textbox(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.2))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_paragraph(p, item.get("title") or "", 40, TITLE_COLOR, bold=True)

    subtitle = (item.get("subtitle") or "").strip()
    if subtitle:
        tf2 = add_textbox(slide, Inches(0.8), Inches(4.5), Inches(11.7), Inches(1.4))
        style_paragraph(tf2.paragraphs[0], subtitle, 22, MUTED)


def render_section(slide, item: dict) -> None:
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8),
        Inches(3.3),
        Inches(1.1),
        Inches(0.12),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    tf = add_textbox(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(2.0))
    style_paragraph(tf.paragraphs[0], item.get("title") or "", 36, TITLE_COLOR, bold=True)


def render_bullets(slide, item: dict) -> None:
    heading = add_textbox(slide, Inches(0.8), Inches(0.55), Inches(11.7), Inches(1.2))
    style_paragraph(heading.paragraphs[0], item.get("title") or "", 28, TITLE_COLOR, bold=True)

    bullets = item.get("bullets") or []
    if not isinstance(bullets, list):
        raise SystemExit("Each bullets slide needs a list of strings in 'bullets'.")
    bullets = [str(b).strip() for b in bullets if str(b).strip()][:8]

    body = add_textbox(slide, Inches(0.9), Inches(1.9), Inches(11.5), Inches(5.0))
    if not bullets:
        return
    for i, bullet in enumerate(bullets):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.level = 0
        para.space_after = Pt(12)
        style_paragraph(para, f"•  {bullet}", 22, BODY)


def load_spec(path: Path) -> dict:
    try:
        spec = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in {path}: {exc}") from exc
    if not isinstance(spec, dict):
        raise SystemExit("slides.json must be an object.")
    if spec.get("intent") not in VALID_INTENTS:
        raise SystemExit(
            "intent must be follow_instructions, teaching_outline, or mixed."
        )
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise SystemExit("slides.json needs a non-empty 'slides' array.")
    for i, item in enumerate(slides):
        if not isinstance(item, dict):
            raise SystemExit(f"slides[{i}] must be an object.")
        if item.get("layout") not in VALID_LAYOUTS:
            raise SystemExit(
                f"slides[{i}].layout must be title, section, or bullets."
            )
        if not str(item.get("title") or "").strip():
            raise SystemExit(f"slides[{i}] needs a title.")
    return spec


def build(spec: dict, dest: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for item in spec["slides"]:
        slide = prs.slides.add_slide(blank)
        fill_slide(slide, BG)
        layout = item["layout"]
        if layout == "title":
            render_title(slide, item)
        elif layout == "section":
            render_section(slide, item)
        else:
            render_bullets(slide, item)
        notes_for(slide, item.get("notes"))

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))


def main() -> None:
    parser = argparse.ArgumentParser(description="Build pptx from slides.json")
    parser.add_argument("spec", type=Path, help="Path to slides.json")
    parser.add_argument("output", type=Path, help="Output .pptx path")
    args = parser.parse_args()
    if not args.spec.is_file():
        raise SystemExit(f"Missing spec file: {args.spec}")
    spec = load_spec(args.spec)
    build(spec, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
